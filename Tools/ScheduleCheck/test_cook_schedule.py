#!/usr/bin/env python3
"""Cross-check the schedule cook against Cory's Week 32 first-look workbook.

The xlsx is not in git. Set SCHEDULE_XLSX, or place the file at the upload
path used in development. When the file is absent the test skips so CI does
not fail closed on a private workbook.
"""

import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(__file__))
import cook_schedule as cook


class QualifyTests(unittest.TestCase):
    def test_blank_week_is_not_zero(self):
        self.assertEqual(cook.measured_schedule(None, None, 100.0, 0.0, 0.0), (None, None, None))
        self.assertEqual(cook.measured_schedule(4.0, 2.0, 90.0, None, None), (4.0, 2.0, 90.0))

    def test_gate(self):
        self.assertFalse(cook.qualifies(None, 20, 20, 20))
        self.assertFalse(cook.qualifies(29999.99, 20, 20, 20))
        self.assertTrue(cook.qualifies(30000, 10, 0, 0))
        self.assertFalse(cook.qualifies(30000, 9.99, 9, 14.99))
        self.assertFalse(cook.qualifies(30000, 0, 9, 0))
        self.assertTrue(cook.qualifies(30000, 0, 9.01, 0))
        self.assertTrue(cook.qualifies(30000, 0, 0, 15))
        self.assertFalse(cook.qualifies(30000, None, None, None))


class WorkbookCrossCheckTests(unittest.TestCase):
    def test_week32_summary_math(self):
        path = os.environ.get("SCHEDULE_XLSX") or "/home/ubuntu/.cursor/projects/workspace/uploads/sched_e2a3.xlsx"
        if not path or not os.path.isfile(path):
            self.skipTest("Schedule Review workbook is not on this machine")
        pack = cook.cook_workbook(path)
        report = pack["crossCheck"]
        self.assertEqual(report["week"], 32)
        self.assertEqual(report["scope"], 2163)
        self.assertEqual(report["underCount"], 2008)
        self.assertEqual(report["overCount"], 1566)
        self.assertEqual(cook.round2(report["eff"]), 64.97)
        self.assertEqual(cook.round2(report["pch"]), 70.28)
        self.assertEqual(cook.round2(report["marketUnder"]), 41.07)
        self.assertEqual(cook.round2(report["marketOver"]), 4.03)
        self.assertEqual(cook.round2(report["storeUnder"]), 23.14)
        self.assertEqual(cook.round2(report["storeOver"]), 13.08)
        self.assertNotEqual(cook.round2(report["storeUnder"]), 41.07)
        self.assertEqual(report["actionCount"], 472)
        self.assertEqual(report["bannerCount"], 468)
        self.assertNotEqual(report["actionCount"], report["bannerCount"])
        self.assertEqual(
            report["regions"],
            {
                "East Region": 612,
                "South Region": 397,
                "California Region": 600,
                "West Region": 554,
            },
        )
        self.assertIn("Week 31", report["summaryTitle"])
        self.assertNotIn("Week 32", report["summaryTitle"])
        shaws = next(item for item in pack["markets"] if item["label"] == "Shaws")
        self.assertEqual(cook.round2(shaws["under"]), 87.99)
        united = next(item for item in pack["markets"] if item["label"] == "United")
        self.assertIsNone(united["under"])
        self.assertIsNone(united["over"])
        store117 = next(item for item in pack["stores"] if item["store"] == "117")
        self.assertEqual(store117["division"], "Shaws")
        self.assertEqual(store117["region"], "East Region")
        self.assertAlmostEqual(store117["sales"], 33961.9775, places=2)
        self.assertGreaterEqual(store117["under"], 99.5)
        self.assertAlmostEqual(store117["eff"], 0.0, places=4)


class CookArgsTests(unittest.TestCase):
    def sample_report(self, **overrides):
        report = {
            "scope": 2163,
            "underCount": 2008,
            "overCount": 1566,
            "eff": 64.97,
            "pch": 70.28,
            "marketUnder": 41.07,
            "marketOver": 4.03,
            "storeUnder": 23.14,
            "storeOver": 13.08,
            "actionCount": 472,
            "bannerCount": 468,
            "notScheduled": 12,
            "regions": {},
            "summaryTitle": "Schedule Review Summary — Week 31",
            "week": 32,
        }
        report.update(overrides)
        return report

    def test_find_positional_is_the_output_path(self):
        args = cook.parse_args(["--find", "/tmp/Heartbeat_Reports", "/tmp/schedule-check.json"])
        self.assertIsNone(args.workbook)
        self.assertEqual(args.output, "/tmp/schedule-check.json")
        self.assertEqual(args.find, "/tmp/Heartbeat_Reports")

    def test_workbook_then_output_stays_positional(self):
        args = cook.parse_args(["/tmp/book.xlsx", "/tmp/schedule-check.json", "--check"])
        self.assertEqual(args.workbook, "/tmp/book.xlsx")
        self.assertEqual(args.output, "/tmp/schedule-check.json")
        self.assertTrue(args.check)

    def test_informational_notes_do_not_fail_company_numbers(self):
        report = self.sample_report()
        self.assertTrue(cook.company_numbers_ok(report))
        self.assertTrue(cook.print_cross_check(report))

    def test_company_number_miss_fails(self):
        report = self.sample_report(scope=1)
        self.assertFalse(cook.company_numbers_ok(report))
        self.assertFalse(cook.print_cross_check(report))

    def test_sqlite_rows_land_in_current_sqlite(self):
        import sqlite3
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            dest = os.path.join(folder, "current.sqlite")
            connection = sqlite3.connect(dest)
            connection.execute("CREATE TABLE facts (id TEXT PRIMARY KEY)")
            connection.execute("INSERT INTO facts(id) VALUES ('keep-me')")
            connection.commit()
            connection.close()
            pack = {
                "publishedAt": "2026-09-28T12:00:00Z",
                "week": 32,
                "filename": "Schedule Review Week 32.xlsx",
                "summaryTitle": "Summary",
                "workbookActionBanner": 468,
                "markets": [{"label": "Total", "under": 41.07, "over": 4.03, "eff": 64.97}],
                "stores": [{
                    "store": "117",
                    "region": "East Region",
                    "division": "Shaws",
                    "district": "B5",
                    "om": "Pat",
                    "sales": 33961.98,
                    "under": 100.0,
                    "over": 0.0,
                    "eff": 0.0,
                    "pch": None,
                    "fourUnder": 9.2,
                    "fourOver": None,
                    "star": None,
                    "dayUnder": [1, None, None, None, None, None, None],
                    "dayOver": [None] * 7,
                }],
                "crossCheck": self.sample_report(),
            }
            original = cook.cook_workbook
            cook.cook_workbook = lambda path, previous_stores=None, **kwargs: pack
            try:
                code = cook.main(["/tmp/unused.xlsx", "--sqlite", dest])
            finally:
                cook.cook_workbook = original
            self.assertEqual(code, 0)
            connection = sqlite3.connect(dest)
            self.assertEqual(connection.execute("SELECT id FROM facts").fetchone()[0], "keep-me")
            row = connection.execute(
                "SELECT week, filename, workbook_action_banner FROM schedule_pack WHERE id = 1"
            ).fetchone()
            self.assertEqual(row, (32, "Schedule Review Week 32.xlsx", 468))
            market = connection.execute("SELECT label, under FROM schedule_market").fetchone()
            self.assertEqual(market[0], "Total")
            self.assertAlmostEqual(market[1], 41.07, places=2)
            store = connection.execute(
                "SELECT store, division, under, day_under_json FROM schedule_store"
            ).fetchone()
            self.assertEqual(store[0], "117")
            self.assertEqual(store[1], "Shaws")
            self.assertAlmostEqual(store[2], 100.0, places=2)
            self.assertIn("1", store[3])
            connection.close()

    def test_publish_sheet_writes_workbook_rows_when_lock_misses(self):
        import sqlite3
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            dest = os.path.join(folder, "current.sqlite")
            sqlite3.connect(dest).close()
            pack = {
                "publishedAt": "2026-10-05T00:00:00Z",
                "week": 32,
                "filename": "Schedule Review Week 32 - Summary.xlsx",
                "summaryTitle": "Schedule Review Summary — Week 31 (WK31)",
                "workbookActionBanner": 468,
                "markets": [{"label": "Total", "under": 5.01, "over": 6.55, "eff": 64.97}],
                "stores": [{
                    "store": "117",
                    "region": "East Region",
                    "division": "Shaws",
                    "district": "B5",
                    "om": "Pat",
                    "sales": 33961.98,
                    "under": 5.01,
                    "over": 6.55,
                    "eff": 64.97,
                    "pch": 70.67,
                    "fourUnder": None,
                    "fourOver": None,
                    "star": None,
                    "dayUnder": [None] * 7,
                    "dayOver": [None] * 7,
                }],
                "crossCheck": self.sample_report(marketUnder=5.01, marketOver=6.55, pch=70.67),
            }
            original = cook.cook_workbook
            cook.cook_workbook = lambda path, previous_stores=None, **kwargs: pack
            try:
                code = cook.main(["/tmp/unused.xlsx", "--sqlite", dest, "--publish-sheet"])
            finally:
                cook.cook_workbook = original
            self.assertEqual(code, 0)
            connection = sqlite3.connect(dest)
            market = connection.execute("SELECT under, over FROM schedule_market WHERE label = 'Total'").fetchone()
            self.assertAlmostEqual(market[0], 5.01, places=2)
            self.assertAlmostEqual(market[1], 6.55, places=2)
            self.assertNotAlmostEqual(market[0], 41.07, places=2)
            connection.close()

    def test_publish_sheet_refuses_empty_store_list(self):
        import sqlite3
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            dest = os.path.join(folder, "current.sqlite")
            sqlite3.connect(dest).close()
            pack = {"week": 32, "stores": [], "crossCheck": self.sample_report(scope=1)}
            original = cook.cook_workbook
            cook.cook_workbook = lambda path, previous_stores=None, **kwargs: pack
            try:
                code = cook.main(["/tmp/unused.xlsx", "--sqlite", dest, "--publish-sheet"])
            finally:
                cook.cook_workbook = original
            self.assertEqual(code, 1)
            connection = sqlite3.connect(dest)
            tables = {
                row[0]
                for row in connection.execute("SELECT name FROM sqlite_master WHERE type = 'table'")
            }
            self.assertNotIn("schedule_store", tables)
            connection.close()

    def test_company_miss_does_not_write_sqlite(self):
        import sqlite3
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            dest = os.path.join(folder, "current.sqlite")
            sqlite3.connect(dest).close()
            pack = {"week": 32, "stores": [], "crossCheck": self.sample_report(scope=1)}
            original = cook.cook_workbook
            cook.cook_workbook = lambda path, previous_stores=None, **kwargs: pack
            try:
                code = cook.main(["/tmp/unused.xlsx", "--sqlite", dest])
            finally:
                cook.cook_workbook = original
            self.assertEqual(code, 1)
            connection = sqlite3.connect(dest)
            tables = {
                row[0]
                for row in connection.execute("SELECT name FROM sqlite_master WHERE type = 'table'")
            }
            self.assertNotIn("schedule_store", tables)
            connection.close()

    def test_matching_numbers_write_json(self):
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            out = os.path.join(folder, "schedule-check.json")
            pack = {"week": 32, "stores": [], "crossCheck": self.sample_report()}
            original = cook.cook_workbook
            cook.cook_workbook = lambda path, previous_stores=None, **kwargs: pack
            try:
                code = cook.main(["/tmp/unused.xlsx", out])
            finally:
                cook.cook_workbook = original
            self.assertEqual(code, 0)
            self.assertTrue(os.path.isfile(out))

    def test_company_miss_removes_json_and_exits_nonzero(self):
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            out = os.path.join(folder, "schedule-check.json")
            pack = {"week": 32, "stores": [], "crossCheck": self.sample_report(scope=1)}
            original = cook.cook_workbook
            cook.cook_workbook = lambda path, previous_stores=None, **kwargs: pack
            try:
                code = cook.main(["/tmp/unused.xlsx", out])
            finally:
                cook.cook_workbook = original
            self.assertEqual(code, 1)
            self.assertFalse(os.path.isfile(out))

    def test_find_with_no_workbook_exits_nonzero(self):
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            out = os.path.join(folder, "schedule-check.json")
            code = cook.main(["--find", folder, out])
            self.assertEqual(code, 1)
            self.assertFalse(os.path.isfile(out))

    def test_missing_path_exits_nonzero(self):
        code = cook.main([])
        self.assertEqual(code, 2)


class WeekFooterTests(unittest.TestCase):
    def write_workbook(
        self,
        folder,
        *,
        week_name=34,
        footer="202634",
        current_footer=None,
        omit_current_footer=False,
        title=None,
        calc=None,
        k1=None,
        over=0.06,
        under=0.04,
        eff=0.80,
    ):
        import openpyxl

        path = os.path.join(folder, f"Schedule Review Week {week_name} - Summary.xlsx")
        book = openpyxl.Workbook()
        current = book.active
        current.title = "Stores Current Week"
        current.append(["Division", "District", "Store", "Over", "Under", "Eff"])
        current.append(["Shaws", "B5", 117, over, under, eff])
        if not omit_current_footer:
            current_id = footer if current_footer is None else current_footer
            current.append([f"Applied filters: WEEK_ID is {current_id}"])

        for name in (
            "Last 4 Week Quality",
            "Sales AVG Last 4 Wks",
            "5 Star Last 5 Weeks",
            "Roster",
            "ACTION NEEDED",
        ):
            book.create_sheet(name)
        for name in ("Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"):
            book.create_sheet(name)

        market = book.create_sheet("Market Look WK33")
        market.append(["Market", "Under", "Over", "Eff"])
        market.append(["Total", 0.05, 0.04, 0.91])
        market.append([f"Applied filters: WEEK_ID is {footer}"])

        store_look = book.create_sheet("Store Look WK33")
        store_look.append([f"Applied filters: WEEK_ID is {footer}"])

        summary = book.create_sheet("Summary")
        shown = title or f"Schedule Review Summary — Week {week_name} (WK{week_name})"
        summary["C1"] = shown
        if k1 is not None:
            summary["K1"] = k1

        if calc is not None:
            hidden = book.create_sheet("_Calc")
            hidden["M2"] = calc["M2"]
            hidden["N2"] = calc["N2"]
            hidden["O2"] = calc["O2"]
            hidden.sheet_state = "hidden"

        book.save(path)
        return path

    def test_renamed_file_uses_footer_week_not_tab_names(self):
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            path = self.write_workbook(folder)
            pack = cook.cook_workbook(path)
        self.assertEqual(pack["week"], 34)
        self.assertIn("Week 34", pack["filename"])
        self.assertTrue(any(item["label"] == "Total" for item in pack["markets"]))
        store = pack["stores"][0]
        self.assertEqual(store["store"], "117")
        self.assertEqual(store["under"], 4.0)
        self.assertEqual(store["over"], 6.0)
        self.assertEqual(store["eff"], 80.0)

    def test_filename_footer_mismatch_fails(self):
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            path = self.write_workbook(folder, footer="202633")
            with self.assertRaises(SystemExit) as caught:
                cook.cook_workbook(path)
        message = str(caught.exception)
        self.assertIn("33", message)
        self.assertIn("34", message)
        self.assertIn("WEEK_ID", message)

    def published_sqlite(self, folder, week):
        import sqlite3

        dest = os.path.join(folder, "current.sqlite")
        connection = sqlite3.connect(dest)
        connection.executescript(cook.SCHEDULE_DDL)
        connection.execute(
            """
            INSERT INTO schedule_pack(id, published_at, week, filename, summary_title)
            VALUES (1, '', ?, 'Schedule Review Week published.xlsx', '')
            """,
            (week,),
        )
        connection.execute(
            """
            INSERT INTO schedule_store(
                store, region, division, district, om, under, over, eff, day_under_json, day_over_json
            ) VALUES ('117', '', '', '', '', 4.0, 6.0, 80.0, '[]', '[]')
            """
        )
        connection.commit()
        connection.close()
        return dest

    def test_stale_identical_store_rows_fail(self):
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            path = self.write_workbook(folder)
            dest = self.published_sqlite(folder, 33)
            with self.assertRaises(SystemExit) as caught:
                cook.main([path, "--sqlite", dest, "--publish-sheet"])
        self.assertIn("byte-identical", str(caught.exception))
        self.assertIn("previous published week", str(caught.exception))

    def test_snapshot_of_previous_week_rejects_stale_rows(self):
        import json
        import sqlite3
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            path = self.write_workbook(folder)
            dest = os.path.join(folder, "current.sqlite")
            sqlite3.connect(dest).close()
            snap_dir = os.path.join(folder, "schedule-weeks")
            os.makedirs(snap_dir)
            with open(os.path.join(snap_dir, "33.json"), "w", encoding="utf-8") as handle:
                json.dump(
                    {"week": 33, "stores": [{"store": "117", "under": 4.0, "over": 6.0, "eff": 80.0}]},
                    handle,
                )
            with self.assertRaises(SystemExit) as caught:
                cook.main([path, "--sqlite", dest, "--publish-sheet"])
        self.assertIn("byte-identical", str(caught.exception))

    def test_changed_store_rows_are_not_stale(self):
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            path = self.write_workbook(folder)
            pack = cook.cook_workbook(
                path,
                previous_stores=[{"store": "117", "under": 4.0, "over": 6.0, "eff": 70.0}],
            )
        self.assertEqual(pack["week"], 34)

    def test_current_week_footer_disagrees(self):
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            path = self.write_workbook(folder, current_footer="202633")
            with self.assertRaises(SystemExit) as caught:
                cook.cook_workbook(path)
        self.assertIn("Stores Current Week", str(caught.exception))
        self.assertIn("202633", str(caught.exception))

    def test_calc_sheet_agreement_passes_as_footer_week(self):
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            path = self.write_workbook(folder, calc={"M2": 34, "N2": 34, "O2": 34})
            pack = cook.cook_workbook(path)
        self.assertEqual(pack["week"], 34)

    def test_calc_sheet_disagreement_fails(self):
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            path = self.write_workbook(folder, calc={"M2": 34, "N2": 33, "O2": 34})
            with self.assertRaises(SystemExit) as caught:
                cook.cook_workbook(path)
        message = str(caught.exception)
        self.assertIn("disagreement", message)
        self.assertIn("N2 33", message)
        self.assertIn("footer 34", message)

    def test_summary_k1_warning_fails(self):
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            path = self.write_workbook(
                folder,
                calc={"M2": 34, "N2": 34, "O2": 34},
                k1="Filename week does not match the Market Look week",
            )
            with self.assertRaises(SystemExit) as caught:
                cook.cook_workbook(path)
        self.assertIn("Summary K1", str(caught.exception))

    def test_two_market_look_sheets_fail(self):
        import openpyxl
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            path = self.write_workbook(folder)
            book = openpyxl.load_workbook(path)
            book.create_sheet("Market Look WK34")
            book.save(path)
            with self.assertRaises(SystemExit) as caught:
                cook.cook_workbook(path)
        self.assertIn("Market Look", str(caught.exception))
        self.assertNotIn("Market Look*", str(caught.exception))

    def test_market_look_notes_is_not_a_look_sheet(self):
        import openpyxl
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            path = self.write_workbook(folder)
            book = openpyxl.load_workbook(path)
            book.create_sheet("Market Look Notes")
            book.save(path)
            pack = cook.cook_workbook(path)
        self.assertEqual(pack["week"], 34)

    def test_whitespace_only_k1_fails(self):
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            path = self.write_workbook(folder, k1="   ")
            with self.assertRaises(SystemExit) as caught:
                cook.cook_workbook(path)
        self.assertIn("Summary K1", str(caught.exception))

    def test_missing_current_week_footer_is_logged(self):
        import io
        import tempfile
        from contextlib import redirect_stdout

        with tempfile.TemporaryDirectory() as folder:
            path = self.write_workbook(folder, omit_current_footer=True)
            buffer = io.StringIO()
            with redirect_stdout(buffer):
                pack = cook.cook_workbook(path)
        self.assertEqual(pack["week"], 34)
        self.assertIn("Stores Current Week has no WEEK_ID footer", buffer.getvalue())

    def test_fallback_without_calc_still_fails_on_k1(self):
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            path = self.write_workbook(folder, k1="Filename week does not match the Market Look week")
            with self.assertRaises(SystemExit) as caught:
                cook.cook_workbook(path)
        self.assertIn("Summary K1", str(caught.exception))

    def test_year_rollover_week_id(self):
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            path = self.write_workbook(folder, week_name=1, footer="202701")
            pack = cook.cook_workbook(path)
        self.assertEqual(pack["week"], 1)
        self.assertIn("Week 1", pack["filename"])

    def test_same_workbook_recook_passes(self):
        import io
        import tempfile
        from contextlib import redirect_stdout

        with tempfile.TemporaryDirectory() as folder:
            path = self.write_workbook(folder)
            dest = self.published_sqlite(folder, 34)
            buffer = io.StringIO()
            with redirect_stdout(buffer):
                code = cook.main([path, "--sqlite", dest, "--publish-sheet"])
        self.assertEqual(code, 0)
        self.assertIn("Stale-row check skipped", buffer.getvalue())
        self.assertNotIn("byte-identical", buffer.getvalue())

    def test_blank_calc_shows_stale_current_week_first(self):
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            path = self.write_workbook(
                folder,
                current_footer="202633",
                calc={"M2": None, "N2": None, "O2": None},
            )
            with self.assertRaises(SystemExit) as caught:
                cook.cook_workbook(path)
        message = str(caught.exception)
        self.assertIn("Stores Current Week", message)
        self.assertIn("202633", message)
        self.assertNotIn("blank", message)


if __name__ == "__main__":
    unittest.main()
